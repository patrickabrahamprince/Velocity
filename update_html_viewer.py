import openpyxl
import json

wb = openpyxl.load_workbook('Velocity_Instagram_All_Comments_Scraped.xlsx')
data = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if rows:
        data[name] = {
            'headers': [str(h or '') for h in rows[0]],
            'rows': [[str(c or '') for c in r] for r in rows[1:]]
        }

json_payload = json.dumps(data)

html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Velocity Circuit - 180+ Instagram Comments & Reels Live Viewer</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #0a0a0a; color: #f3f4f6; margin: 0; padding: 24px; }}
  h1 {{ color: #CCFF00; font-size: 24px; margin-bottom: 8px; }}
  p {{ color: #9ca3af; margin-top: 0; margin-bottom: 24px; font-size: 14px; }}
  .tabs {{ display: flex; gap: 8px; margin-bottom: 16px; border-bottom: 1px solid #27272a; padding-bottom: 8px; }}
  .tab {{ background: #18181b; border: 1px solid #27272a; color: #d4d4d8; padding: 8px 16px; border-radius: 8px; cursor: pointer; font-weight: 600; font-size: 13px; transition: all 0.2s; }}
  .tab.active {{ background: #CCFF00; color: #000; border-color: #CCFF00; }}
  .search-box {{ width: 100%; max-width: 450px; padding: 10px 14px; background: #18181b; border: 1px solid #3f3f46; border-radius: 8px; color: white; margin-bottom: 16px; font-size: 14px; }}
  .table-container {{ overflow-x: auto; background: #121214; border: 1px solid #27272a; border-radius: 12px; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  th {{ background: #18181b; color: #a1a1aa; text-align: left; padding: 12px 14px; border-bottom: 1px solid #27272a; position: sticky; top: 0; font-weight: 600; text-transform: uppercase; font-size: 11px; letter-spacing: 0.5px; }}
  td {{ padding: 12px 14px; border-bottom: 1px solid #1f1f23; color: #e4e4e7; vertical-align: top; }}
  tr:hover {{ background: rgba(204, 255, 0, 0.04); }}
  a {{ color: #60a5fa; text-decoration: none; }}
  a:hover {{ text-decoration: underline; }}
  .badge {{ display: inline-block; padding: 2px 8px; border-radius: 999px; font-size: 11px; font-weight: 600; background: #27272a; color: #CCFF00; border: 1px solid #3f3f46; }}
</style>
</head>
<body>
  <h1>🏎️ Velocity International Circuit — Full Instagram Comments Dataset</h1>
  <p>180+ real comments extracted across all reels with auto-intent detection & recommended responses.</p>
  
  <div class="tabs" id="tabs"></div>
  <input type="text" id="search" class="search-box" placeholder="🔍 Search comments, handles, topics, queries..." oninput="filterTable()" />
  
  <div class="table-container" id="tableContainer"></div>

  <script>
    const data = {json_payload};
    let currentSheet = Object.keys(data)[0];

    function renderTabs() {{
      const tabsEl = document.getElementById('tabs');
      tabsEl.innerHTML = Object.keys(data).map(name => 
        `<button class="tab ${{name === currentSheet ? 'active' : ''}}" onclick="switchSheet('${{name}}')">${{name.replace(/_/g, ' ')}}</button>`
      ).join('');
    }}

    function switchSheet(name) {{
      currentSheet = name;
      renderTabs();
      renderTable();
    }}

    function renderTable() {{
      const sheet = data[currentSheet];
      const query = document.getElementById('search').value.toLowerCase();
      
      const filteredRows = sheet.rows.filter(row => 
        row.some(cell => cell.toLowerCase().includes(query))
      );

      let html = '<table><thead><tr>' + sheet.headers.map(h => `<th>${{h}}</th>`).join('') + '</tr></thead><tbody>';
      
      if (filteredRows.length === 0) {{
        html += `<tr><td colspan="${{sheet.headers.length}}" style="text-align:center; padding: 32px; color: #71717a;">No matching records found</td></tr>`;
      }} else {{
        filteredRows.forEach(row => {{
          html += '<tr>' + row.map((cell, idx) => {{
            if (cell.startsWith('http')) {{
              return `<td><a href="${{cell}}" target="_blank">${{cell}}</a></td>`;
            }}
            if (cell.startsWith('@')) {{
              return `<td><span class="badge">${{cell}}</span></td>`;
            }}
            return `<td>${{cell}}</td>`;
          }}).join('') + '</tr>';
        }});
      }}
      
      html += '</tbody></table>';
      document.getElementById('tableContainer').innerHTML = html;
    }}

    function filterTable() {{
      renderTable();
    }}

    renderTabs();
    renderTable();
  </script>
</body>
</html>
"""

with open('instagram_viewer.html', 'w', encoding='utf-8') as f:
    f.write(html_content)

print("Updated instagram_viewer.html successfully")
